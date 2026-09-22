# -*- coding: utf-8 -*-
"""将 PRD_source.md 组装为规范排版的 Word 文档（.docx）。

支持：标题 1-4 级 / 段落 / 加粗 / 项目符号与编号 / 表格 / 引用块 / 图片嵌入 /
封面页 / 静态目录页 / 页眉页脚页码。
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, "PRD_source.md")
OUT = os.path.join(BASE, "PRD_英语学习网站MVP.docx")

HEADING_EAST = {1: "黑体", 2: "黑体", 3: "黑体", 4: "微软雅黑"}
HEADING_SIZE = {1: 16, 2: 14, 3: 12, 4: 11}
HEADING_COLOR = {1: (0x1F, 0x4E, 0x79), 2: (0x1F, 0x4E, 0x79),
                 3: (0x33, 0x33, 0x33), 4: (0x33, 0x33, 0x33)}


def set_font(run, size=10.5, bold=False, italic=False, color=None,
             latin="Times New Roman", east="宋体"):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), east)


def add_rich(par, text, size=10.5, base_bold=False):
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for i, part in enumerate(parts):
        if part == "":
            continue
        run = par.add_run(part)
        set_font(run, size=size, bold=base_bold or (i % 2 == 1))


def para(doc, text, size=10.5, space_after=6, align=None, indent=None,
         bold=False, color=None, italic=False):
    p = doc.add_paragraph()
    add_rich(p, text, size=size, base_bold=bold)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    if color:
        for r in p.runs:
            r.font.color.rgb = RGBColor(*color)
    if italic:
        for r in p.runs:
            r.font.italic = True
    return p


def quote_para(doc, text):
    p = doc.add_paragraph()
    add_rich(p, text, size=10)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.55)
    pf.space_after = Pt(8)
    for r in p.runs:
        r.font.italic = True
        r.font.color.rgb = RGBColor(0x55, 0x63, 0x70)
    return p


def add_image(doc, alt, path):
    is_proto = "prototype" in path
    caption = alt
    if is_proto and not alt.startswith("原型"):
        caption = "原型线框图：" + alt
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(15.0))
    p.paragraph_format.space_after = Pt(2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_font(r, size=9, color=((0x1F, 0x4E, 0x79) if is_proto else (0x66, 0x66, 0x66)))
    cap.paragraph_format.space_after = Pt(10)


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = "Table Grid"
    tbl.alignment = 1  # CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = tbl.cell(i, j)
            cell.paragraphs[0].text = ""
            text = row[j] if j < len(row) else ""
            add_rich(cell.paragraphs[0], text, size=10)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            if i == 0:
                for r in cell.paragraphs[0].runs:
                    set_font(r, size=10, bold=True, color=(0x1F, 0x4E, 0x79))
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "EAF1F8")
                cell._element.get_or_add_tcPr().append(shd)
    return tbl


def style_setup(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for lvl in (1, 2, 3, 4):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(HEADING_SIZE[lvl])
        st.font.bold = True
        st.font.color.rgb = RGBColor(*HEADING_COLOR[lvl])
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_EAST[lvl])
        st.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        st.paragraph_format.space_after = Pt(8)


def add_page_number(doc):
    sec = doc.sections[0]
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run("第 ")
    set_font(r1, size=9)
    field = p.add_run("")
    set_font(field, size=9)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    field._r.append(fld1)
    field._r.append(instr)
    field._r.append(fld2)
    r2 = p.add_run(" 页")
    set_font(r2, size=9)
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("英语学习网站 · MVP 产品需求文档 v0.1")
    set_font(hr, size=9, color=(0x99, 0x99, 0x99))


def cover_page(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = para(doc, "英语学习网站", size=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.runs[0].text = "英语学习网站"
    set_font(p.runs[0], size=30, bold=True, east="黑体")
    p2 = para(doc, "产品需求文档（PRD）", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p2.runs[0], size=18, bold=True, east="微软雅黑")
    p3 = para(doc, "MVP 版本 v0.1", align=WD_ALIGN_PARAGRAPH.CENTER,
              color=(0x66, 0x66, 0x66))
    set_font(p3.runs[0], size=12)
    for _ in range(8):
        doc.add_paragraph()
    meta = [
        "文档状态：MVP 定稿（已评审）",
        "创建日期：2026-09-20",
        "适用范围：Web 端（桌面 + 手机浏览器，响应式）",
        "北极星指标：7 日留存",
    ]
    for line in meta:
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER,
             color=(0x44, 0x44, 0x44))
    doc.add_page_break()


def toc_page(doc, headings):
    para(doc, "目录", size=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.paragraphs[-1].runs[0].text = "目录"
    set_font(doc.paragraphs[-1].runs[0], size=18, bold=True, east="黑体")
    for lvl, text in headings:
        indent_map = {1: 0.0, 2: 0.6, 3: 1.2}
        ind = indent_map.get(lvl, 1.2)
        if lvl == 1:
            p = doc.add_paragraph()
            add_rich(p, text, size=12, base_bold=True)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph()
            add_rich(p, text, size=10.5)
            p.paragraph_format.left_indent = Cm(ind)
            p.paragraph_format.space_after = Pt(1)
    doc.add_page_break()


def parse_lines():
    with open(SRC, encoding="utf-8") as f:
        lines = [ln.rstrip("\n") for ln in f]
    headings = []
    body = []
    for ln in lines:
        if ln.startswith("#"):
            m = re.match(r"^(#{1,4})\s+(.*)$", ln)
            if m:
                headings.append((len(m.group(1)), m.group(2)))
                body.append(("heading", len(m.group(1)), m.group(2)))
                continue
        body.append(("raw", ln))
    return headings, body


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(25)
    sec.bottom_margin = Mm(20)
    sec.left_margin = Mm(25)
    sec.right_margin = Mm(25)
    sec.different_first_page_header_footer = True
    style_setup(doc)
    add_page_number(doc)

    prop = doc.core_properties
    prop.title = "英语学习网站 MVP 产品需求文档（PRD）"
    prop.author = "产品组"
    prop.subject = "英语学习网站 MVP"

    headings, body = parse_lines()
    cover_page(doc)
    toc_page(doc, headings)

    i = 0
    n = len(body)
    while i < n:
        kind = body[i][0]
        if kind == "heading":
            lvl, text = body[i][1], body[i][2]
            if not text.startswith("英语学习网站 MVP 产品需求文档"):
                h = doc.add_heading("", level=lvl)
                add_rich(h, text, size=HEADING_SIZE[lvl], base_bold=True)
            i += 1
            continue
        ln = body[i][1]
        if ln.strip() == "":
            i += 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < n and body[i][1].lstrip().startswith("|"):
                raw = body[i][1].strip()
                cells = [c.strip() for c in raw.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            doc.add_paragraph()
            continue
        img = re.match(r"^!\[(.+?)\]\((.+?)\)$", ln)
        if img:
            path = os.path.join(BASE, img.group(2))
            if os.path.exists(path):
                add_image(doc, img.group(1), path)
            else:
                para(doc, f"[缺失图片：{img.group(1)}]",
                     color=(0xB3, 0x35, 0x2B))
            i += 1
            continue
        if ln.startswith(">"):
            quote_para(doc, ln.lstrip("> ").strip())
            i += 1
            continue
        m = re.match(r"^(\s*)-\s+(.*)$", ln)
        if m:
            depth = 0 if m.group(1) == "" else 1
            style = "List Bullet" if depth == 0 else ("List Bullet 2"
                                                      if depth == 1 else None)
            p = doc.add_paragraph(style=style)
            add_rich(p, m.group(2))
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue
        m = re.match(r"^(\s*)\d+\.\s+(.*)$", ln)
        if m:
            style = "List Number" if m.group(1) == "" else None
            p = doc.add_paragraph(style=style)
            add_rich(p, m.group(2))
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue
        para(doc, ln)
        i += 1

    doc.save(OUT)
    print("generated:", OUT)


if __name__ == "__main__":
    build()