# -*- coding: utf-8 -*-
"""验证生成的 PRD .docx：结构完整性 / 图片嵌入 / 表格 / 关键章节 / 可打开性。"""
import os
from zipfile import ZipFile
from docx import Document

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(BASE, "PRD_英语学习网站MVP.docx")
KEYWORDS = ["项目背景", "目标用户", "功能范围", "用户故事", "里程碑",
            "非功能需求", "首学漏斗", "每日学习总流程", "缺席分批容错",
            "间隔调度", "北极星指标", "修订记录", "术语表"]

fails = []

with ZipFile(DOCX) as z:
    names = z.namelist()
    assert "word/document.xml" in names, "缺少 document.xml"
    media = [n for n in names if n.startswith("word/media/")]
    print(f"[OK] zip 结构完整，内嵌图片 {len(media)} 张: {media}")

doc = Document(DOCX)
paras = [p.text for p in doc.paragraphs]
full = "\n".join(paras)
tables = doc.tables
imgs = len(media)

print(f"[OK] 段落数={len(paras)}，表格数={len(tables)}，图片数={imgs}")

missing = [k for k in KEYWORDS if k not in full]
if missing:
    fails.append(f"缺少关键词: {missing}")
else:
    print(f"[OK] 关键章节齐全（{len(KEYWORDS)} 项全部命中）")

headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
print(f"[OK] 标题总数={len(headings)}：")
for h in headings[:28]:
    print("    -", h)

if imgs < 6:
    fails.append(f"图片数量不足，期望>=6，实际={imgs}")
if len(tables) < 15:
    fails.append(f"表格数量偏少，期望>=15，实际={len(tables)}")
if len(paras) < 200:
    fails.append(f"段落过少，期望>=200，实际={len(paras)}")

print()
if fails:
    for f in fails:
        print("[FAIL]", f)
    raise SystemExit(1)
print("[PASS] 全部结构校验通过")